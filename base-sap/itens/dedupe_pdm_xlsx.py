"""Desambigua descrições UC duplicadas (mesma família + pdm_signature) no xlsx SAP."""
from __future__ import annotations

import re
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl

XLSX = Path(__file__).with_name("Base de itens SAP B1.xlsx")
SHEET = "Uso e consumo"


def pdm_signature(raw: str) -> str:
    text = unicodedata.normalize("NFD", raw or "")
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    text = text.upper()
    text = "".join(
        c
        for c in text
        if not unicodedata.category(c).startswith("P")
        and not unicodedata.category(c).startswith("S")
    )
    return re.sub(r"\s+", " ", text).strip()


def resolve_family(family: str, subgroup: str) -> str:
    fam = (family or "").upper().strip()
    sub = (subgroup or "").upper().strip()
    if fam == "MATERIAL ADESIVO E VEDACAO" and sub == "MANUTENCAO E OBRAS":
        return "MANUTENCAO E OBRAS"
    return fam


def main() -> None:
    wb = openpyxl.load_workbook(XLSX)
    ws = wb[SHEET]
    rows: list[tuple[int, str, str, str]] = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, max_col=8, values_only=True), start=2):
        sap = str(row[0] or "").strip().upper()
        desc = str(row[1] or "").strip()
        family = resolve_family(str(row[7] or ""), str(row[6] or ""))
        if not sap or not desc:
            continue
        rows.append((idx, sap, desc, family))

    groups: dict[tuple[str, str], list[tuple[int, str, str]]] = defaultdict(list)
    for idx, sap, desc, family in rows:
        groups[(family, pdm_signature(desc))].append((idx, sap, desc))

    changed = 0
    log: list[str] = []
    for (family, sig), items in groups.items():
        if len(items) < 2:
            continue
        items.sort(key=lambda x: x[1])
        keep = items[0]
        for idx, sap, desc in items[1:]:
            suffix = f" {sap}"
            if desc.upper().endswith(sap):
                continue
            new_desc = f"{desc.strip()}{suffix}"
            ws.cell(row=idx, column=2).value = new_desc
            changed += 1
            log.append(f"{sap}\t{family}\t{keep[1]} kept\t{desc} -> {new_desc}")

    wb.save(XLSX)
    print(f"Grupos com colisao: {sum(1 for v in groups.values() if len(v) > 1)}")
    print(f"Descricoes ajustadas: {changed}")
    for line in log:
        print(line)

    # re-read verify
    wb2 = openpyxl.load_workbook(XLSX, read_only=True)
    ws2 = wb2[SHEET]
    again: dict[tuple[str, str], list[str]] = defaultdict(list)
    for row in ws2.iter_rows(min_row=2, max_col=8, values_only=True):
        sap = str(row[0] or "").strip().upper()
        desc = str(row[1] or "").strip()
        family = resolve_family(str(row[7] or ""), str(row[6] or ""))
        if not sap or not desc:
            continue
        again[(family, pdm_signature(desc))].append(sap)
    leftover = {k: v for k, v in again.items() if len(v) > 1}
    print(f"Colisoes restantes: {len(leftover)}")
    for k, v in leftover.items():
        print(k, v)


if __name__ == "__main__":
    main()
